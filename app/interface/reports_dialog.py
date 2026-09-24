"""Reportes de caja: ventas, ganancias y gastos por rango de fechas.

Ventana configurable con rango de fechas, modo detalle/agrupado y una gráfica
(QtCharts) que acompaña los datos.
"""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime, time, timedelta
from decimal import Decimal

from PySide6.QtCharts import (
    QBarCategoryAxis,
    QBarSeries,
    QBarSet,
    QChart,
    QChartView,
    QPieSeries,
    QValueAxis,
)
from PySide6.QtCore import Qt
from PySide6.QtGui import QColor, QPainter, QTextDocument
from PySide6.QtPrintSupport import QPrinter
from PySide6.QtWidgets import (
    QCheckBox,
    QComboBox,
    QDateEdit,
    QDialog,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QListWidget,
    QListWidgetItem,
    QPushButton,
    QSplitter,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from app.application.bus import QueryBus
from app.application.queries import GetCashMovementsQuery, GetCatalogQuery, GetSalesHistoryQuery
from app.domain.value_objects import Money
from app.interface.dialogs import MOVEMENT_LABELS
from app.interface.terminal_colors import terminal_label
from app.interface.widgets import make_label, make_table

SALE_COMPLETED = "COMPLETADA"
CASH_OUT = "SALIDA"


def _terminal_cell_color(receipt_number: str) -> "object":
    """Color de fondo suave según la caja del ticket (para pintar celdas)."""
    from app.interface.terminal_colors import terminal_row_color

    return terminal_row_color(receipt_number)

REPORTS = (
    ("ventas", "Ventas"),
    ("ganancias", "Ganancias"),
    ("gastos", "Gastos"),
)

AGG_KEYS = {
    "ventas": (("dia", "Día"), ("producto", "Producto"), ("categoria", "Categoría"), ("metodo", "Método de pago")),
    "ganancias": (("dia", "Día"), ("producto", "Producto"), ("categoria", "Categoría")),
    "gastos": (("dia", "Día"), ("concepto", "Concepto")),
}

KEY_HEADERS = {
    "dia": "Fecha",
    "producto": "Producto",
    "categoria": "Categoría",
    "metodo": "Método",
    "concepto": "Concepto",
}

CHART_BARS = "Barras"
CHART_PIE = "Pastel"

PALETTE = ("#ef4444", "#2563eb", "#16a34a", "#f59e0b", "#8b5cf6", "#06b6d4", "#ec4899",
           "#84cc16", "#f97316", "#0ea5e9", "#14b8a6", "#f43f5e")

MAX_CHART_SLICES = 20


def _money_divide(money: Money, divisor: int) -> Money:
    return Money(money.amount / Decimal(divisor)) if divisor else Money.zero()


def _fmt_money_signed(value: Decimal) -> str:
    sign = "-" if value < 0 else ""
    return sign + Money(abs(value)).format()


def _money_pct(part: Decimal, whole: Decimal) -> str:
    if whole <= 0:
        return "—"
    return f"{part / whole * 100:.1f}%"


class ReportsDialog(QDialog):
    """Ventana de reportes configurables por rango de fechas y agrupación."""

    def __init__(self, queries: QueryBus, parent=None):
        super().__init__(parent)
        self._queries = queries
        self._info: dict[int, dict] = {}
        self._export_title = ""
        self._export_headers: list[str] = []
        self._export_rows: list[list[str]] = []
        self._export_totals: list[str] = []
        self.setWindowTitle("Reportes")
        self.resize(1000, 640)

        root = QVBoxLayout(self)
        root.setContentsMargins(14, 14, 14, 14)
        root.setSpacing(10)
        root.addWidget(make_label("Reportes", object_name="pageTitle"))

        body = QHBoxLayout()
        body.setSpacing(12)

        self.report_list = QListWidget()
        self.report_list.setFixedWidth(170)
        for code, label in REPORTS:
            item = QListWidgetItem(label)
            item.setData(Qt.UserRole, code)
            self.report_list.addItem(item)
        self.report_list.setCurrentRow(0)

        right = QVBoxLayout()
        right.setSpacing(10)

        filters = QHBoxLayout()
        self.preset = QComboBox()
        self.preset.addItems(
            ["Hoy", "Ayer", "Esta semana", "Este mes", "Últimos 7 días", "Últimos 30 días", "Todo"]
        )
        self.start_date = QDateEdit(calendarPopup=True)
        self.start_date.setDisplayFormat("dd/MM/yyyy")
        self.end_date = QDateEdit(calendarPopup=True)
        self.end_date.setDisplayFormat("dd/MM/yyyy")
        filters.addWidget(self.preset)
        filters.addSpacing(10)
        filters.addWidget(make_label("Desde:", object_name="muted"))
        filters.addWidget(self.start_date)
        filters.addWidget(make_label("Hasta:", object_name="muted"))
        filters.addWidget(self.end_date)
        filters.addStretch(1)
        generate_btn = QPushButton("Generar")
        generate_btn.setObjectName("primary")
        generate_btn.clicked.connect(self._generate)
        filters.addWidget(generate_btn)
        right.addLayout(filters)

        options = QHBoxLayout()
        self.mode = QComboBox()
        self.mode.addItem("Detalle", "detalle")
        self.mode.addItem("Agrupado", "agrupado")
        self.agroup = QComboBox()
        self.show_chart = QCheckBox("Mostrar gráfica")
        self.show_chart.setChecked(True)
        self.chart_type = QComboBox()
        self.chart_type.addItems([CHART_BARS, CHART_PIE])
        options.addWidget(make_label("Modo:", object_name="muted"))
        options.addWidget(self.mode)
        options.addSpacing(10)
        options.addWidget(make_label("Agrupar por:", object_name="muted"))
        options.addWidget(self.agroup)
        options.addSpacing(10)
        options.addWidget(self.show_chart)
        options.addWidget(make_label("Gráfica:", object_name="muted"))
        options.addWidget(self.chart_type)
        options.addStretch(1)
        right.addLayout(options)

        self.summary = make_label("", object_name="muted")
        right.addWidget(self.summary)

        splitter = QSplitter(Qt.Horizontal)
        self.table = make_table(["Recibo", "Fecha", "Ítems", "Total", "Método"], stretch_column=0)
        splitter.addWidget(self.table)

        self.chart_frame = QFrame()
        self.chart_frame.setObjectName("card")
        chart_layout = QVBoxLayout(self.chart_frame)
        chart_layout.setContentsMargins(8, 8, 8, 8)
        self.chart_holder = QVBoxLayout()
        chart_layout.addLayout(self.chart_holder)
        splitter.addWidget(self.chart_frame)
        splitter.setStretchFactor(0, 3)
        splitter.setStretchFactor(1, 2)
        splitter.setSizes([600, 380])
        right.addWidget(splitter, 1)

        body.addWidget(self.report_list)
        body.addLayout(right, 1)
        root.addLayout(body, 1)

        footer = QHBoxLayout()
        export_pdf_btn = QPushButton("Exportar PDF")
        export_excel_btn = QPushButton("Exportar Excel")
        close_btn = QPushButton("Cerrar")
        close_btn.setObjectName("primary")
        for btn in (export_pdf_btn, export_excel_btn):
            btn.setObjectName("ghost")
        export_pdf_btn.clicked.connect(self._export_pdf)
        export_excel_btn.clicked.connect(self._export_excel)
        footer.addWidget(export_pdf_btn)
        footer.addWidget(export_excel_btn)
        footer.addStretch(1)
        footer.addWidget(close_btn)
        root.addLayout(footer)

        self._refresh_agroup()
        self._apply_preset("Este mes")

        self.report_list.currentRowChanged.connect(lambda _: self._on_options_changed())
        self.mode.currentIndexChanged.connect(lambda _: self._on_options_changed())
        self.agroup.currentIndexChanged.connect(lambda _: self._on_options_changed())
        self.show_chart.toggled.connect(lambda _: self._generate())
        self.chart_type.currentIndexChanged.connect(lambda _: self._generate())
        self.start_date.dateChanged.connect(lambda _: self._generate())
        self.end_date.dateChanged.connect(lambda _: self._generate())
        self.preset.currentIndexChanged.connect(self._on_preset)

        self._generate()

    # ------------------------------------------------------------------ #
    # Configuración de rangos
    # ------------------------------------------------------------------ #

    def _range(self):
        start = self.start_date.date().toPython()
        end = self.end_date.date().toPython()
        return datetime.combine(start, time.min), datetime.combine(end, time.max)

    def _on_preset(self, index: int) -> None:
        self._apply_preset(self.preset.currentText())

    def _apply_preset(self, name: str) -> None:
        today = datetime.now().date()
        self.start_date.blockSignals(True)
        self.end_date.blockSignals(True)
        if name == "Hoy":
            start = end = today
        elif name == "Ayer":
            start = end = today - timedelta(days=1)
        elif name == "Esta semana":
            start = today - timedelta(days=today.weekday())
            end = today
        elif name == "Este mes":
            start = today.replace(day=1)
            end = today
        elif name == "Últimos 7 días":
            start = today - timedelta(days=6)
            end = today
        elif name == "Últimos 30 días":
            start = today - timedelta(days=29)
            end = today
        else:  # Todo
            start = today.replace(year=today.year - 10, month=1, day=1)
            end = today
        self.start_date.setDate(self._to_qdate(start))
        self.end_date.setDate(self._to_qdate(end))
        self.start_date.blockSignals(False)
        self.end_date.blockSignals(False)

    @staticmethod
    def _to_qdate(date) -> "object":
        from PySide6.QtCore import QDate

        return QDate(date.year, date.month, date.day)

    # ------------------------------------------------------------------ #
    # Estado
    # ------------------------------------------------------------------ #

    def _current_report(self) -> str:
        item = self.report_list.currentItem()
        return str(item.data(Qt.UserRole)) if item else "ventas"

    def _on_options_changed(self) -> None:
        self._refresh_agroup()
        self._generate()

    def _refresh_agroup(self) -> None:
        report = self._current_report()
        keys = AGG_KEYS.get(report, AGG_KEYS["ventas"])
        current = self.agroup.currentData()
        self.agroup.blockSignals(True)
        self.agroup.clear()
        for code, label in keys:
            self.agroup.addItem(label, code)
        index = self.agroup.findData(current if current in {c for c, _ in keys} else "dia")
        self.agroup.setCurrentIndex(max(0, index))
        self.agroup.blockSignals(False)

    def _label_header(self, key: str) -> str:
        return KEY_HEADERS.get(key, "Período")

    # ------------------------------------------------------------------ #
    # Generación
    # ------------------------------------------------------------------ #

    def _generate(self) -> None:
        report = self._current_report()
        mode = str(self.mode.currentData())
        key = str(self.agroup.currentData())
        chart_key = "dia" if mode == "detalle" else key
        start, end = self._range()

        if report == "ventas":
            sales = [
                s for s in self._queries.ask(GetSalesHistoryQuery(start=start, end=end))
                if s.status == SALE_COMPLETED
            ]
            self._render_sales(sales, mode, key, chart_key)
        elif report == "ganancias":
            sales = [
                s for s in self._queries.ask(GetSalesHistoryQuery(start=start, end=end))
                if s.status == SALE_COMPLETED
            ]
            catalog = self._queries.ask(GetCatalogQuery(include_inactive=True))
            self._info = {
                p.id: {"cost": p.cost if p.cost is not None else Money.zero(), "category": p.category_name}
                for p in catalog
            }
            self._render_profit(sales, mode, key, chart_key)
        else:
            movements = [
                m for m in self._queries.ask(GetCashMovementsQuery(start=start, end=end))
                if m.movement_type == CASH_OUT
            ]
            self._render_expenses(movements, mode, key, chart_key)

    # ------------------------------------------------------------------ #
    # Ventas
    # ------------------------------------------------------------------ #

    def _render_sales(self, sales, mode: str, key: str, chart_key: str) -> None:
        total = sum((s.total for s in sales), Money.zero())
        n = len(sales)
        avg = _money_divide(total, n) if n else Money.zero()
        self.summary.setText(f"{n} ventas · Total: {total.format()} · Ticket promedio: {avg.format()}")

        if mode == "detalle":
            headers = ["Recibo", "Caja", "Fecha", "Ítems", "Total", "Método"]
            items_total = sum(s.item_count for s in sales)
            rows = [
                [s.receipt_number, terminal_label(s.receipt_number),
                 s.created_at.strftime("%d/%m/%Y %H:%M"),
                 str(s.item_count), s.total.format(), s.methods_label]
                for s in sales
            ]
            stretch, right_align = 0, {4}
            totals = ["TOTALES", "", "", str(items_total), total.format(), ""]
        else:
            agg = self._aggregate_sales(sales, key, self._sale_parts_sales)
            headers = [self._label_header(key), "Ventas", "Total"]
            rows = [[row["label"], str(row["count"]), row["amount"].format()] for row in agg]
            stretch, right_align = 0, {2}
            totals = ["TOTALES", str(n), total.format()]

        chart = self._chart_series_sales(sales, chart_key)
        self._render(headers, rows, totals, stretch, right_align,
                     chart, "Ventas por período", color_columns={0, 1})

    def _aggregate_sales(self, sales, key: str, parts) -> list[dict]:
        buckets: dict[str, dict] = defaultdict(lambda: {"count": 0, "amount": Money.zero(), "cost": Money.zero()})
        for sale in sales:
            for label, amount, cost, count in parts(sale, key):
                bucket = buckets[label]
                bucket["count"] += count
                bucket["amount"] = bucket["amount"] + amount
                bucket["cost"] = bucket["cost"] + cost
        return self._finalize_agg(buckets, key)

    def _chart_series_sales(self, sales, key: str):
        agg = self._aggregate_sales(sales, key, self._sale_parts_sales)
        return agg, [("Ventas", [row["amount"].amount for row in agg])]

    def _sale_parts_sales(self, sale, key: str):
        day = sale.created_at.strftime("%d/%m/%Y")
        if key == "dia":
            return [(day, sale.total, sale.total, 1)]
        if key == "metodo":
            return [(p.method.value, p.amount, p.amount, 1) for p in sale.payments]
        parts = []
        for item in sale.items:
            qty = item.quantity - item.refunded_qty
            if qty <= 0:
                continue
            if key == "producto":
                label = item.product_name
            else:
                info = self._info.get(item.product_id, {})
                label = info.get("category") or "Sin categoría"
            parts.append((label, item.unit_price * qty, item.unit_price * qty, qty))
        return parts

    # ------------------------------------------------------------------ #
    # Ganancias
    # ------------------------------------------------------------------ #

    def _render_profit(self, sales, mode: str, key: str, chart_key: str) -> None:
        income = sum((self._sale_totals(s, self._info)[0] for s in sales), Money.zero())
        cost = sum((self._sale_totals(s, self._info)[1] for s in sales), Money.zero())
        profit = income.amount - cost.amount
        self.summary.setText(
            f"Ingresos: {income.format()} · Costo: {cost.format()} · "
            f"Ganancia: {_fmt_money_signed(profit)} ({_money_pct(profit, income.amount)})"
        )

        if mode == "detalle":
            headers = ["Recibo", "Caja", "Fecha", "Ingresos", "Costo", "Ganancia", "Margen"]
            rows = []
            for sale in sales:
                sale_income, sale_cost = self._sale_totals(sale, self._info)
                sale_profit = sale_income.amount - sale_cost.amount
                rows.append([
                    sale.receipt_number, terminal_label(sale.receipt_number),
                    sale.created_at.strftime("%d/%m/%Y %H:%M"),
                    sale_income.format(), sale_cost.format(), _fmt_money_signed(sale_profit),
                    _money_pct(sale_profit, sale_income.amount),
                ])
            stretch, right_align = 0, {3, 4, 5}
            totals = ["TOTALES", "", "", income.format(), cost.format(), _fmt_money_signed(profit), ""]
        else:
            agg = self._aggregate_sales(sales, key, self._sale_parts_profit)
            headers = [self._label_header(key), "Ventas", "Ingresos", "Costo", "Ganancia"]
            rows = [
                [row["label"], str(row["count"]), row["amount"].format(),
                 row["cost"].format(),
                 _fmt_money_signed(row["amount"].amount - row["cost"].amount)]
                for row in agg
            ]
            stretch, right_align = 0, {2, 3, 4}
            totals = ["TOTALES", str(len(sales)), income.format(), cost.format(), _fmt_money_signed(profit)]

        chart = self._chart_series_profit(sales, chart_key)
        self._render(headers, rows, totals, stretch, right_align, chart, "Ganancias por período",
                     color_columns={0, 1})

    def _chart_series_profit(self, sales, key: str):
        agg = self._aggregate_sales(sales, key, self._sale_parts_profit)
        if key == "dia":
            series = [
                ("Ingresos", [row["amount"].amount for row in agg]),
                ("Costo", [row["cost"].amount for row in agg]),
                ("Ganancia", [row["amount"].amount - row["cost"].amount for row in agg]),
            ]
        else:
            series = [("Ganancia", [row["amount"].amount - row["cost"].amount for row in agg])]
        return agg, series

    def _sale_parts_profit(self, sale, key: str):
        if key == "dia":
            income, cost = self._sale_totals(sale, self._info)
            return [(sale.created_at.strftime("%d/%m/%Y"), income, cost, 1)]
        parts = []
        for item in sale.items:
            qty = item.quantity - item.refunded_qty
            if qty <= 0:
                continue
            info = self._info.get(item.product_id, {})
            cost = info.get("cost", Money.zero()) * qty
            if key == "producto":
                label = item.product_name
            else:
                label = info.get("category") or "Sin categoría"
            parts.append((label, item.unit_price * qty, cost, qty))
        return parts

    def _sale_totals(self, sale, info: dict[int, dict]):
        income = Money.zero()
        cost = Money.zero()
        for item in sale.items:
            qty = item.quantity - item.refunded_qty
            if qty <= 0:
                continue
            income = income + item.unit_price * qty
            inf = info.get(item.product_id, {})
            cost = cost + (inf.get("cost", Money.zero()) * qty)
        income = income - sale.discount
        return income, cost

    # ------------------------------------------------------------------ #
    # Gastos
    # ------------------------------------------------------------------ #

    def _render_expenses(self, movements, mode: str, key: str, chart_key: str) -> None:
        total = sum((m.amount for m in movements), Money.zero())
        self.summary.setText(f"{len(movements)} egresos · Total: {total.format()}")

        if mode == "detalle":
            headers = ["Fecha", "Concepto", "Nota", "Monto"]
            rows = [
                [m.created_at.strftime("%d/%m/%Y %H:%M"),
                 MOVEMENT_LABELS.get(m.reason, m.reason or "—"), m.note or "", m.amount.format()]
                for m in movements
            ]
            stretch, right_align = 2, {3}
            totals = ["TOTALES", "", "", total.format()]
        else:
            agg = self._aggregate_movements(movements, key)
            headers = [self._label_header(key), "Movimientos", "Total"]
            rows = [[row["label"], str(row["count"]), row["amount"].format()] for row in agg]
            stretch, right_align = 0, {2}
            totals = ["TOTALES", str(len(movements)), total.format()]

        agg, series = self._chart_series_expenses(movements, chart_key)
        self._render(headers, rows, totals, stretch, right_align, (agg, series), "Egresos por período")

    def _aggregate_movements(self, movements, key: str) -> list[dict]:
        buckets: dict[str, dict] = defaultdict(lambda: {"count": 0, "amount": Money.zero()})
        for movement in movements:
            if key == "dia":
                label = movement.created_at.strftime("%d/%m/%Y")
            else:
                label = MOVEMENT_LABELS.get(movement.reason, movement.reason or "—")
            bucket = buckets[label]
            bucket["count"] += 1
            bucket["amount"] = bucket["amount"] + movement.amount
        return self._finalize_agg(buckets, key)

    def _chart_series_expenses(self, movements, key: str):
        agg = self._aggregate_movements(movements, key)
        return agg, [("Egresos", [row["amount"].amount for row in agg])]

    # ------------------------------------------------------------------ #
    # Utilidades
    # ------------------------------------------------------------------ #

    @staticmethod
    def _finalize_agg(buckets: dict, key: str) -> list[dict]:
        rows = [dict(label=label, **values) for label, values in buckets.items()]
        if key == "dia":
            rows.sort(key=lambda r: datetime.strptime(r["label"], "%d/%m/%Y"))
        else:
            rows.sort(key=lambda r: r["amount"].amount, reverse=True)
        return rows

    def _render(self, headers, rows, totals, stretch: int, right_align, chart_data, chart_title: str, *,
            color_columns: set[int] | None = None) -> None:
        self._export_title = chart_title
        self._export_headers = list(headers)
        self._export_rows = [list(row) for row in rows]
        self._export_totals = list(totals)

        self.table.setColumnCount(len(headers))
        self.table.setHorizontalHeaderLabels(headers)
        self.table.setRowCount(0)
        header = self.table.horizontalHeader()
        header.setStretchLastSection(False)
        for column in range(len(headers)):
            header.setSectionResizeMode(column, QHeaderView.ResizeToContents)
        if 0 <= stretch < len(headers):
            header.setSectionResizeMode(stretch, QHeaderView.Stretch)
        for row in rows:
            r = self.table.rowCount()
            self.table.insertRow(r)
            fallback_bg = _terminal_cell_color(str(row[0])) if row else None
            for column, value in enumerate(row):
                item = QTableWidgetItem(str(value))
                if column in right_align:
                    item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                if color_columns and column in color_columns and fallback_bg is not None:
                    item.setBackground(fallback_bg)
                self.table.setItem(r, column, item)
        self._append_totals_row(totals, right_align)
        agg, series = chart_data
        self._show_chart([row["label"] for row in agg], series, chart_title)

    def _append_totals_row(self, totals, right_align) -> None:
        from PySide6.QtGui import QBrush, QColor

        if not totals or not any(totals):
            return
        r = self.table.rowCount()
        self.table.insertRow(r)
        for column, value in enumerate(totals):
            item = QTableWidgetItem(str(value))
            font = item.font()
            font.setBold(True)
            item.setFont(font)
            brush = QBrush(QColor("#eef2ff"))
            item.setBackground(brush)
            if column in right_align:
                item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
            self.table.setItem(r, column, item)

    # ------------------------------------------------------------------ #
    # Exportación
    # ------------------------------------------------------------------ #

    def _export_pdf(self) -> None:
        target, _ = QFileDialog.getSaveFileName(
            self, "Exportar reporte a PDF", "reporte.pdf", "PDF (*.pdf)"
        )
        if not target:
            return
        if not target.lower().endswith(".pdf"):
            target += ".pdf"
        try:
            doc = QTextDocument()
            doc.setHtml(self._export_html())
            printer = QPrinter(QPrinter.HighResolution)
            printer.setOutputFormat(QPrinter.PdfFormat)
            printer.setOutputFileName(target)
            doc.print_(printer)
        except Exception as exc:  # noqa: BLE001
            from PySide6.QtWidgets import QMessageBox

            QMessageBox.critical(self, "No se pudo exportar", str(exc))
            return
        from PySide6.QtWidgets import QMessageBox

        QMessageBox.information(self, "Exportado", f"Reporte guardado en:\n{target}")

    def _export_excel(self) -> None:
        target, _ = QFileDialog.getSaveFileName(
            self, "Exportar reporte a Excel", "reporte.xlsx", "Libro Excel (*.xlsx)"
        )
        if not target:
            return
        if not target.lower().endswith(".xlsx"):
            target += ".xlsx"
        try:
            self._write_excel(target)
        except Exception as exc:  # noqa: BLE001
            from PySide6.QtWidgets import QMessageBox

            QMessageBox.critical(self, "No se pudo exportar", str(exc))
            return
        from PySide6.QtWidgets import QMessageBox

        QMessageBox.information(self, "Exportado", f"Reporte guardado en:\n{target}")

    def _export_html(self) -> str:
        """HTML tabular del reporte actual, listo para imprimir/exportar a PDF."""
        from html import escape

        lines = ["<html><head><meta charset='utf-8'></head><body>"]
        lines.append(f"<h2>{escape(self._export_title)}</h2>")
        lines.append(f"<p>{escape(self.summary.text())}</p>")
        lines.append("<table border='1' cellpadding='5' cellspacing='0' "
                     "style='border-collapse:collapse;font-family:sans-serif;font-size:12px'>")
        lines.append("<tr>" + "".join(f"<th style='background:#f3f4f6'>{escape(h)}</th>" for h in self._export_headers) + "</tr>")
        for row in self._export_rows:
            lines.append("<tr>" + "".join(f"<td>{escape(str(c))}</td>" for c in row) + "</tr>")
        if self._export_totals and any(self._export_totals):
            lines.append("<tr>" + "".join(
                f"<td style='font-weight:bold;background:#eef2ff'>{escape(str(c))}</td>"
                for c in self._export_totals
            ) + "</tr>")
        lines.append("</table></body></html>")
        return "".join(lines)

    def _write_excel(self, target) -> None:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte"
        ws.cell(row=1, column=1, value=self._export_title).font = Font(bold=True, size=14)
        ws.cell(row=2, column=1, value=self.summary.text()).font = Font(italic=True)
        header_row = 4
        for col, header in enumerate(self._export_headers, start=1):
            cell = ws.cell(row=header_row, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="F3F4F6")
        row = header_row + 1
        for data in self._export_rows:
            for col, value in enumerate(data, start=1):
                ws.cell(row=row, column=col, value=value)
            row += 1
        if self._export_totals and any(self._export_totals):
            for col, value in enumerate(self._export_totals, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="EEF2FF")
                if col != 1:
                    cell.alignment = Alignment(horizontal="right")
        for col in range(1, len(self._export_headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20
        wb.save(target)

    # ------------------------------------------------------------------ #
    # Gráficas
    # ------------------------------------------------------------------ #

    def _limit_chart(self, categories, series):
        """Limita la gráfica a las N categorías con mayor valor absoluto."""
        if len(categories) <= MAX_CHART_SLICES:
            return categories, series
        primary = series[0][1]
        indices = sorted(
            range(len(categories)),
            key=lambda i: abs(float(primary[i])),
            reverse=True,
        )[:MAX_CHART_SLICES]
        indices.sort()
        limited_categories = [categories[i] for i in indices]
        limited_series = [(name, [values[i] for i in indices]) for name, values in series]
        return limited_categories, limited_series

    def _show_chart(self, categories, series, title: str) -> None:
        while self.chart_holder.count():
            item = self.chart_holder.takeAt(0)
            widget = item.widget()
            if widget is not None:
                widget.deleteLater()

        has_data = bool(categories) and bool(series) and any(
            any(v != 0 for v in values) for _, values in series
        )
        if not self.show_chart.isChecked():
            self.chart_frame.setVisible(False)
            return
        if not has_data:
            self.chart_frame.setVisible(True)
            self.chart_holder.addWidget(make_label("Sin datos en el período.", object_name="muted", alignment=Qt.AlignCenter))
            return
        self.chart_frame.setVisible(True)
        categories, series = self._limit_chart(categories, series)
        chart = self._make_chart(categories, series, str(self.chart_type.currentText()), title)
        if chart is None:
            return
        view = QChartView(chart)
        view.setRenderHint(QPainter.Antialiasing)
        self.chart_holder.addWidget(view)

    def _make_chart(self, categories, series, chart_type: str, title: str):
        chart = QChart()
        chart.setTitle(title)
        legend = chart.legend()
        legend.setVisible(True)
        legend.setAlignment(Qt.AlignBottom)

        if chart_type == CHART_PIE:
            name, values = series[0]
            pie = QPieSeries()
            for index, (cat, value) in enumerate(zip(categories, values)):
                if value <= 0:
                    continue
                piece = pie.append(f"{cat}", float(value))
                piece.setBrush(QColor(PALETTE[index % len(PALETTE)]))
            if not pie.slices():
                return None
            chart.addSeries(pie)
            for piece in pie.slices():
                piece.setLabelVisible(True)
                piece.setLabel(f"{piece.label()} {piece.value():,.0f}")
        else:
            bar = QBarSeries()
            for index, (name, values) in enumerate(series):
                bar_set = QBarSet(name)
                bar_set.setColor(QColor(PALETTE[index % len(PALETTE)]))
                for value in values:
                    bar_set.append(float(value))
                bar.append(bar_set)
            chart.addSeries(bar)
            axis_x = QBarCategoryAxis()
            axis_x.append(list(categories))
            chart.addAxis(axis_x, Qt.AlignBottom)
            bar.attachAxis(axis_x)
            axis_y = QValueAxis()
            axis_y.setLabelFormat("%.1f")
            chart.addAxis(axis_y, Qt.AlignLeft)
            bar.attachAxis(axis_y)
        return chart