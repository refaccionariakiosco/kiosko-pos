"""Exportador de la lista de compra a un libro Excel (.xlsx)."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from app.application.read_models import PurchaseLineDTO
from app.domain.value_objects import Money
from app.settings import StoreInfo


def write_purchase_list_excel(
    lines: list[PurchaseLineDTO],
    path: str | Path,
    store: StoreInfo | None = None,
) -> Path:
    """Escribe el pedido en un .xlsx con encabezado, detalle y total."""
    import openpyxl
    from openpyxl.styles import Alignment, Font

    target = Path(path)
    store = store or StoreInfo()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lista de compra"

    headers = ["Proveedor", "Código", "Descripción", "Precio unit.", "Cant.", "Subtotal"]
    widths = [22, 14, 40, 14, 9, 14]

    row = 1
    if store.name:
        cell = ws.cell(row=row, column=1, value=store.name)
        cell.font = Font(bold=True, size=14)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        row += 1
    date_cell = ws.cell(row=row, column=1, value=f"Lista de compra · {datetime.now():%d/%m/%Y %H:%M}")
    date_cell.font = Font(italic=True)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
    row += 2

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True)
        ws.column_dimensions[ws.cell(row=row, column=col).column_letter].width = widths[col - 1]
    row += 1

    total = Money.zero()
    for line in lines:
        subtotal = line.unit_price * line.quantity
        total += subtotal
        ws.cell(row=row, column=1, value=line.provider_name)
        ws.cell(row=row, column=2, value=line.code)
        ws.cell(row=row, column=3, value=line.description)
        price_cell = ws.cell(row=row, column=4, value=line.unit_price.amount)
        price_cell.number_format = "0.00"
        ws.cell(row=row, column=5, value=line.quantity)
        sub_cell = ws.cell(row=row, column=6, value=subtotal.amount)
        sub_cell.number_format = "0.00"
        row += 1

    total_cell = ws.cell(row=row, column=5, value="TOTAL")
    total_cell.font = Font(bold=True)
    total_cell.alignment = Alignment(horizontal="right")
    amount_cell = ws.cell(row=row, column=6, value=total.amount)
    amount_cell.font = Font(bold=True)
    amount_cell.number_format = "0.00"

    wb.save(target)
    return target