"""Pruebas del área de proveedores (independiente del inventario)."""

from __future__ import annotations

import zipfile

import pytest

from app.application import (
    CreateProviderCommand,
    DeleteProviderCommand,
    ImportProviderItemsCommand,
    ListProviderItemsQuery,
    ListProvidersQuery,
    ProviderItemRequest,
    ReceiveLineRequest,
    ReceivePurchaseOrderCommand,
    UpdateProviderCommand,
)
from app.domain.entities import Provider, ProviderItem
from app.domain.exceptions import ProviderNotFoundError, ValidationError
from app.domain.value_objects import Money
from app.infrastructure.importers.provider_list import (
    ColumnInfo,
    guess_column,
    preview_worksheet,
    read_provider_items_xlsx,
)
from app.infrastructure.printers.receipt import render_purchase_list_html
from app.application.read_models import PurchaseLineDTO


def _make_xlsx(rows: list[list[object]], path: str) -> None:
    """Crea un .xlsx mínimo (sharedStrings + sheet1) con las filas dadas."""
    shared: dict[str, int] = {}
    cells: list[list[tuple[str, str, str]]] = []
    for row in rows:
        row_cells: list[tuple[str, str, str]] = []
        for idx, value in enumerate(row, start=1):
            col = chr(64 + idx)
            if isinstance(value, bool) or value is None:
                continue
            if isinstance(value, (int, float)):
                row_cells.append((col, "n", str(value)))
            else:
                text = str(value)
                if text not in shared:
                    shared[text] = len(shared)
                row_cells.append((col, "s", str(shared[text])))
        cells.append(row_cells)

    shared_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="%d">'
        + "".join(f"<si><t>{text}</t></si>" for text in shared)
        + "</sst>"
    ) % len(shared)
    sheet = ['<?xml version="1.0" encoding="UTF-8"?>']
    sheet.append('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>')
    for r, row_cells in enumerate(cells, start=1):
        inner = "".join(f'<c r="{c}{r}" t="{t}"><v>{v}</v></c>' for c, t, v in row_cells)
        sheet.append(f'<row r="{r}">{inner}</row>')
    sheet.append("</sheetData></worksheet>")

    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        zf.writestr("_rels/.rels", '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>')
        zf.writestr("xl/workbook.xml", '<?xml version="1.0"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheets><sheet name="Sheet1"/></sheets></workbook>')
        zf.writestr("xl/sharedStrings.xml", shared_xml)
        zf.writestr("xl/worksheets/sheet1.xml", "".join(sheet))


# --------------------------------------------------------------------------- #
# Dominio
# --------------------------------------------------------------------------- #


def test_provider_requiere_nombre() -> None:
    with pytest.raises(ValidationError):
        Provider(id=None, name="   ")


def test_provider_item_valida_campos() -> None:
    with pytest.raises(ValidationError):
        ProviderItem(code="", description="Agua", price=Money.from_input("10"))
    with pytest.raises(ValidationError):
        ProviderItem(code="A1", description="", price=Money.from_input("10"))
    with pytest.raises(ValidationError):
        ProviderItem(code="A1", description="Agua", price=Money.zero())


# --------------------------------------------------------------------------- #
# CQRS + persistencia
# --------------------------------------------------------------------------- #


def test_crea_proveedor_e_importa_lista(services) -> None:
    provider = services.commands.execute(CreateProviderCommand(name="Distribuidora 9", phone="555-1234"))
    assert provider.id > 0
    assert provider.name == "Distribuidora 9"

    count = services.commands.execute(
        ImportProviderItemsCommand(
            provider_id=provider.id,
            items=(
                ProviderItemRequest(code="A-1", description="Gaseosa 1.5L", price="1800.50"),
                ProviderItemRequest(code="A-2", description="Agua 500ml", price="700"),
            ),
        )
    )
    assert count == 2

    providers = services.queries.ask(ListProvidersQuery())
    assert len(providers) == 1
    assert providers[0].item_count == 2

    items = services.queries.ask(ListProviderItemsQuery(provider_id=provider.id))
    assert len(items) == 2
    by_code = {item.code: item for item in items}
    assert by_code["A-1"].price.format() == "$ 1.800,50"
    assert by_code["A-1"].provider_name == "Distribuidora 9"


def test_lista_no_afecta_inventario(services) -> None:
    provider = services.commands.execute(CreateProviderCommand(name="Z"))
    services.commands.execute(
        ImportProviderItemsCommand(
            provider_id=provider.id, items=(ProviderItemRequest(code="X", description="Y", price="5"),)
        )
    )
    catalog = services.queries.ask(__import__("app.application", fromlist=["GetCatalogQuery"]).GetCatalogQuery())
    assert catalog == []


def test_importar_reemplaza_lista(services) -> None:
    provider = services.commands.execute(CreateProviderCommand(name="A"))
    services.commands.execute(
        ImportProviderItemsCommand(provider_id=provider.id, items=(ProviderItemRequest(code="1", description="Uno", price="1"),))
    )
    services.commands.execute(
        ImportProviderItemsCommand(provider_id=provider.id, items=(ProviderItemRequest(code="2", description="Dos", price="2"),))
    )
    items = services.queries.ask(ListProviderItemsQuery(provider_id=provider.id))
    assert [i.code for i in items] == ["2"]


def test_importar_a_proveedor_inexistente(services) -> None:
    with pytest.raises(ProviderNotFoundError):
        services.commands.execute(
            ImportProviderItemsCommand(provider_id=9999, items=(ProviderItemRequest(code="1", description="Uno", price="1"),))
        )


def test_editar_proveedor(services) -> None:
    provider = services.commands.execute(CreateProviderCommand(name="Viejo", phone="111", note="a"))
    updated = services.commands.execute(
        UpdateProviderCommand(provider_id=provider.id, name="Nuevo", phone="222", note="b")
    )
    assert updated.name == "Nuevo"
    assert updated.phone == "222"
    assert updated.note == "b"
    providers = services.queries.ask(ListProvidersQuery())
    assert [p.name for p in providers] == ["Nuevo"]


def test_editar_proveedor_inexistente(services) -> None:
    with pytest.raises(ProviderNotFoundError):
        services.commands.execute(UpdateProviderCommand(provider_id=9999, name="X"))


def test_eliminar_proveedor_borra_sus_articulos(services) -> None:
    provider = services.commands.execute(CreateProviderCommand(name="Para borrar"))
    services.commands.execute(
        ImportProviderItemsCommand(
            provider_id=provider.id,
            items=(
                ProviderItemRequest(code="A", description="Uno", price="1"),
                ProviderItemRequest(code="B", description="Dos", price="2"),
            ),
        )
    )
    services.commands.execute(DeleteProviderCommand(provider_id=provider.id))
    assert services.queries.ask(ListProvidersQuery()) == []
    assert services.queries.ask(ListProviderItemsQuery()) == []


def test_eliminar_proveedor_inexistente(services) -> None:
    with pytest.raises(ProviderNotFoundError):
        services.commands.execute(DeleteProviderCommand(provider_id=9999))


# --------------------------------------------------------------------------- #
# Recepción de pedido (afecta inventario)
# --------------------------------------------------------------------------- #


def test_recepcion_de_pedido_suma_stock_y_relaciona(services) -> None:
    from app.application import CreateProductCommand, GetCatalogQuery

    provider = services.commands.execute(CreateProviderCommand(name="Distribuidora"))
    services.commands.execute(
        ImportProviderItemsCommand(
            provider_id=provider.id,
            items=(
                ProviderItemRequest(code="A-1", description="Gaseosa 1.5L", price="1800.50"),
                ProviderItemRequest(code="A-2", description="Agua 500ml", price="700"),
            ),
        )
    )
    product = services.commands.execute(
        CreateProductCommand(code="A-1", name="Gaseosa 1.5L", unit_price="2500", stock=0)
    )
    items = services.queries.ask(ListProviderItemsQuery(provider_id=provider.id))
    by_code = {i.code: i for i in items}

    received = services.commands.execute(
        ReceivePurchaseOrderCommand(
            note="Recepción de prueba",
            lines=(
                ReceiveLineRequest(
                    provider_item_id=by_code["A-1"].id,
                    product_id=product.id,
                    quantity=5,
                    unit_price="1800.50",
                ),
            ),
        )
    )
    assert received == 5

    catalog = services.queries.ask(__import__("app.application", fromlist=["GetCatalogQuery"]).GetCatalogQuery())
    updated = next(p for p in catalog if p.id == product.id)
    assert updated.stock == 5
    assert updated.cost is not None and updated.cost.amount == 1800.50

    items2 = services.queries.ask(ListProviderItemsQuery(provider_id=provider.id))
    by_code2 = {i.code: i for i in items2}
    assert by_code2["A-1"].product_id == product.id
    assert by_code2["A-2"].product_id is None

    movements = services.queries.ask(__import__("app.application", fromlist=["GetStockMovementsQuery"]).GetStockMovementsQuery(product_id=product.id))
    assert movements[0].delta == 5
    assert movements[0].reason == "COMPRA"


def test_recepcion_sin_renglones_rechaza(services) -> None:
    with pytest.raises(ValidationError):
        services.commands.execute(ReceivePurchaseOrderCommand(note="", lines=()))


def test_find_inventory_match() -> None:
    from app.application.read_models import ProductDTO, ProviderItemDTO
    from app.interface.provider_dialogs import find_inventory_match

    products = [
        ProductDTO(id=1, code="A-1", name="Gaseosa 1.5L", unit_price=Money.from_input("10"), stock=0, min_stock=0, active=True),
        ProductDTO(id=2, code="A-2", name="Agua 500ml", unit_price=Money.from_input("8"), stock=0, min_stock=0, active=True),
    ]
    item = ProviderItemDTO(id=7, provider_id=1, code="A-1", description="Gaseosa", price=Money.from_input("9"))
    assert find_inventory_match(item, products).id == 1

    related = ProviderItemDTO(id=8, provider_id=1, code="ZZZ", description="Algo", price=Money.from_input("9"), product_id=2)
    assert find_inventory_match(related, products).id == 2

    sin_match = ProviderItemDTO(id=9, provider_id=1, code="ZZZ", description="Nada que ver", price=Money.from_input("9"))
    assert find_inventory_match(sin_match, products) is None


# --------------------------------------------------------------------------- #
# Import xlsx con mapeo
# --------------------------------------------------------------------------- #


SAMPLE = [
    ["Codigo", "Producto", "Precio de lista"],
    [1001, "Gaseosa Cola 1.5L", "2400"],
    ["1002", "Galletitas Rellenas", "950.50"],
    ["", "Sin codigo", "999"],
    ["1003", "Sin precio", ""],
]


@pytest.fixture()
def xlsx(tmp_path) -> str:
    path = str(tmp_path / "lista_proveedor.xlsx")
    _make_xlsx(SAMPLE, path)
    return path


def test_vista_previa_mapea_columnas(xlsx: str) -> None:
    preview = preview_worksheet(xlsx)
    headers = {c.letter: c.header for c in preview.columns}
    assert headers == {"A": "Codigo", "B": "Producto", "C": "Precio de lista"}
    assert preview.row_count == 4
    assert len(preview.sample) == 4


def test_guess_column() -> None:
    columns = [ColumnInfo("A", "Codigo"), ColumnInfo("B", "Producto"), ColumnInfo("C", "Precio de lista")]
    assert guess_column(columns, code=True) == "A"
    assert guess_column(columns, description=True) == "B"
    assert guess_column(columns, price=True) == "C"


def test_lectura_con_mapeo_omite_invalidos(xlsx: str) -> None:
    rows, skipped = read_provider_items_xlsx(xlsx, code_col="A", description_col="B", price_col="C")
    assert len(rows) == 2
    assert skipped == 2
    by_code = {r.code: r for r in rows}
    assert by_code["1001"].description == "Gaseosa Cola 1.5L"
    assert float(by_code["1002"].price) == 950.50


# --------------------------------------------------------------------------- #
# Ticket
# --------------------------------------------------------------------------- #


def test_ticket_lista_de_compra() -> None:
    from app.settings import StoreInfo

    lines = [
        PurchaseLineDTO(
            code="A-1",
            description="Gaseosa 1.5L",
            provider_name="Distribuidora 9",
            quantity=2,
            unit_price=Money.from_input("1800.50"),
            subtotal=Money.from_input("3601.00"),
        )
    ]
    store = StoreInfo(name="KIOSCO CENTRAL", footer="gracias")
    html = render_purchase_list_html(lines, store)
    assert "KIOSCO CENTRAL" in html
    assert "Gaseosa 1.5L" in html
    assert "Distribuidora 9" in html
    assert "$ 3.601,00" in html


# --------------------------------------------------------------------------- #
# Export a Excel
# --------------------------------------------------------------------------- #


def test_exportar_lista_de_compra_a_excel(tmp_path: pytest.TempPathFactory) -> None:
    from openpyxl import load_workbook

    from app.infrastructure.exporters.purchase_list import write_purchase_list_excel
    from app.settings import StoreInfo

    lines = [
        PurchaseLineDTO(
            code="A-1",
            description="Gaseosa 1.5L",
            provider_name="Distribuidora 9",
            quantity=2,
            unit_price=Money.from_input("1800.50"),
            subtotal=Money.from_input("3601.00"),
        ),
        PurchaseLineDTO(
            code="B-7",
            description="Galletitas",
            provider_name="Distribuidora 9",
            quantity=3,
            unit_price=Money.from_input("950.00"),
            subtotal=Money.from_input("2850.00"),
        ),
    ]
    store = StoreInfo(name="KIOSCO CENTRAL", footer="gracias")
    target = tmp_path / "pedido.xlsx"
    write_purchase_list_excel(lines, target, store)

    wb = load_workbook(str(target))
    ws = wb.active
    values = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
    texts = ["".join(map(str, [c for c in row if c is not None])) for row in values]
    assert any("KIOSCO CENTRAL" in t for t in texts)
    headers = [c for c in values[3] if c is not None]
    assert headers == ["Proveedor", "Código", "Descripción", "Precio unit.", "Cant.", "Subtotal"]
    assert values[4][1] == "A-1"
    assert values[4][2] == "Gaseosa 1.5L"
    assert float(values[4][4]) == 2.0
    assert float(values[4][5]) == 3601.00
    total_row = values[6]
    assert total_row[4] == "TOTAL"
    assert float(total_row[5]) == 6451.00